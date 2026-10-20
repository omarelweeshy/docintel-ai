import io

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.core.errors import ParseError
from app.ingest.chunking import StructureChunker
from app.ingest.structure import Block, parse_blocks, table_rows, txt_blocks


def test_table_rows_merge_headers_and_align_merged_cells() -> None:
    # Shape of the SP 800-40r4 matrix as PyMuPDF extracts it: spacer columns, a two-row header,
    # and header cells that sit one column to the right of their data.
    raw = [
        ["", "Vulnerability", "", "Asset Importance", "", "", ""],
        ["", "Importance", "", "Low", "", "High", ""],
        ["Low", "", "By deadline: 64.7 %", "", "By deadline: 85.0 %", "", ""],
        ["Medium", "", "By deadline: 66.5 %", "", "By deadline: 71.4 %", "", ""],
    ]
    assert table_rows(raw) == [
        "Vulnerability Importance: Low; Asset Importance Low: By deadline: 64.7 %; "
        "High: By deadline: 85.0 %",
        "Vulnerability Importance: Medium; Asset Importance Low: By deadline: 66.5 %; "
        "High: By deadline: 71.4 %",
    ]


def test_table_rows_fill_row_labels_and_join_wrapped_text() -> None:
    raw = [
        ["Element", "Description", "Priority", "Notes"],
        ["GV.OC-03", "Legal and regulatory", "Medium", "R1: include all"],
        ["", "requirements are understood", "", ""],
        ["", "Second item", "High", "N1"],
    ]
    assert table_rows(raw) == [
        "Element: GV.OC-03; Description: Legal and regulatory requirements are understood; "
        "Priority: Medium; Notes: R1: include all",
        "Element: GV.OC-03; Description: Second item; Priority: High; Notes: N1",
    ]


def test_txt_headings_and_paragraphs() -> None:
    blocks = txt_blocks("RUNBOOK\n\nSCOPE\nCovers the API.\n\nAlerts page the primary first.\n")
    assert [(b.kind, b.text) for b in blocks] == [
        ("heading", "RUNBOOK"),
        ("heading", "SCOPE"),
        ("text", "Covers the API."),
        ("text", "Alerts page the primary first."),
    ]


def test_docx_keeps_headings_and_table_rows() -> None:
    doc = DocxDocument()
    doc.add_heading("Retention", level=1)
    doc.add_paragraph("Table 1 lists periods.")
    table = doc.add_table(rows=0, cols=2)
    for row in [("Class", "Period"), ("Payroll", "6 years"), ("Footage", "45 days")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row
    stream = io.BytesIO()
    doc.save(stream)
    blocks, pages = parse_blocks(stream.getvalue(), ".docx", 100000, 10)
    assert pages is None
    assert [(b.kind, b.text) for b in blocks] == [
        ("heading", "Retention"),
        ("text", "Table 1 lists periods."),
        ("table_row", "Class: Payroll; Period: 6 years"),
        ("table_row", "Class: Footage; Period: 45 days"),
    ]


def pdf(pages: list[list[tuple[str, float]]]) -> bytes:
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page(width=595, height=842)
        y = 120.0
        for text, size in lines:
            page.insert_text((72, y), text, fontsize=size)
            y += size * 2.2
        # Running footer and a vertical watermark on every page.
        page.insert_text((250, 820), "Example Corp | Handbook", fontsize=8)
        page.insert_text((20, 600), "Available free of charge", fontsize=8, rotate=90)
    return doc.tobytes()


def test_pdf_headings_running_footer_and_watermark() -> None:
    data = pdf(
        [[("1 Purpose", 16), ("Body text one.", 10)]] * 2 + [[("2 Scope", 16), ("More.", 10)]]
    )
    blocks, pages = parse_blocks(data, ".pdf", 100000, 10)
    assert pages == 3
    texts = [b.text for b in blocks]
    assert "Example Corp | Handbook" not in texts
    assert not any("free of charge" in t for t in texts)
    assert blocks[0] == Block("heading", "1 Purpose", 1, 1)
    assert Block("text", "Body text one.", 1) in blocks


def test_pdf_without_text_layer_still_requires_ocr() -> None:
    doc = pymupdf.open()
    doc.new_page()
    with pytest.raises(ParseError, match="OCR"):
        parse_blocks(doc.tobytes(), ".pdf", 100000, 10)


def test_chunker_respects_headings_pages_and_size() -> None:
    blocks = [
        Block("heading", "1 Policy", 1, 1),
        Block("text", "word " * 50, 1),
        Block("heading", "1.1 Detail", 1, 2),
        Block("text", "alpha " * 30, 1),
        Block("text", "beta " * 30, 2),
        Block("heading", "2 Other", 2, 1),
        Block("table_row", "Class: A; Period: 1 year", 2, table="t"),
    ]
    chunks = StructureChunker(target_words=40, max_words=60).split(blocks)
    assert [c.text.split("\n")[0] for c in chunks] == [
        "1 Policy",
        "1 Policy > 1.1 Detail",
        "1 Policy > 1.1 Detail",
        "2 Other",
    ]
    assert [c.page_number for c in chunks] == [1, 1, 2, 2]
    assert chunks[0].metadata["section_id"] == 1
    assert chunks[1].metadata["section_id"] == chunks[2].metadata["section_id"] == 2
    assert chunks[3].metadata["kinds"] == ["table_row"]
    # The 50-word paragraph is under max_words and stays whole.
    assert chunks[0].text.count("word") == 50


def test_table_chunks_carry_heading_and_caption() -> None:
    from app.ingest.chunking import table_chunks

    blocks = [
        Block("heading", "3.5 Metrics", 24, 2),
        Block("text", "Table 1: Vulnerability Mitigation Time Summary Matrix", 24),
        Block("table_row", "Vulnerability: High; Low: By deadline: 68.6 %", 24, table="p24t0"),
        Block("table_row", "Vulnerability: Critical; Low: By deadline: 81.4 %", 24, table="p24t0"),
        Block("text", "Organizations should frequently update their low-level metrics.", 24),
    ]
    chunks = table_chunks(blocks, start_index=10)
    assert len(chunks) == 1
    assert chunks[0].index == 10 and chunks[0].page_number == 24
    assert chunks[0].text.splitlines() == [
        "3.5 Metrics > Table 1: Vulnerability Mitigation Time Summary Matrix",
        "Vulnerability: High; Low: By deadline: 68.6 %",
        "Vulnerability: Critical; Low: By deadline: 81.4 %",
    ]
